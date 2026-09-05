"""
src/reporting/metrics.py + src/reporting/summary_builder.py +
src/reporting/exporter.py (Adım 7) testleri.

Gerçek 2019/2023 örnek belgeleriyle üretilen TEK bir ComparisonResult
(modül düzeyinde bir kez hesaplanır -- pahalı ingestion/parse/match/diff
adımlarını her testte TEKRARLAMAMAK için), ground_truth.json'daki bilinen
sayılara karşı objektif doğrulanır.
"""

from __future__ import annotations

import json
from io import BytesIO

import pandas as pd
import pytest

from src.analysis.pipeline import compare_documents
from src.config import CHANGE_TYPE_LABELS_TR, SAMPLES_DIR
from src.models import ChangeType
from src.reporting.exporter import export_to_csv, export_to_excel, export_to_html
from src.reporting.html_renderer import render_summary_html
from src.reporting.metrics import (
    BOLUMSUZ_ETIKETI,
    build_detail_dataframe,
    build_section_breakdown,
    build_structural_dataframe,
    build_summary_dataframe,
)
from src.reporting.summary_builder import SummaryStats, build_summary_stats, render_summary_markdown

with open(SAMPLES_DIR / "ground_truth.json", encoding="utf-8") as f:
    _GROUND_TRUTH = json.load(f)

# NEDEN ground_truth.json'daki eski (6 değerli) "change_type" alanı BURADA
# İKİYE AYRIŞTIRILARAK yorumlanır (dosya DEĞİŞTİRİLMEDİ -- bkz. Adım 9 notu,
# tests/test_differ.py'deki AYNI gerekçe): "RENUMBERED"/"MOVED" etiketi
# "içerik AYNI, sadece numara/yer FARKLI" anlamına gelir -- yeni modelde
# content_status=IDENTICAL + ilgili yapısal bayrak=True'ya karşılık gelir.
#
# NEDEN _EXPECTED_STRUCTURAL_COUNTS "change_type" etiketinden DEĞİL,
# DOĞRUDAN madde_no_2019 != madde_no_2023 karşılaştırmasından türetilir:
# classifier.py::is_renumbered de TAM OLARAK böyle çalışır (etiketten
# BAĞIMSIZ, saf numara karşılaştırması) -- örn. "MODIFIED" etiketli
# birim_sorumluluklari'nın (9->8) numarası DA değişmiştir, sadece "RENUMBERED"
# etiketli maddelere bakmak bu satırı KAÇIRIRDI (bkz. Adım 9'un TAM konusu).
_EXPECTED_COUNTS: dict[ChangeType, int] = dict.fromkeys(ChangeType, 0)
_EXPECTED_STRUCTURAL_COUNTS = {"RENUMBERED": 0, "MOVED": 0}
for _m in _GROUND_TRUTH["maddeler"]:
    _content_type = (
        ChangeType.IDENTICAL
        if _m["change_type"] in ("IDENTICAL", "RENUMBERED", "MOVED")
        else ChangeType(_m["change_type"])
    )
    _EXPECTED_COUNTS[_content_type] += 1
    _no19, _no23 = _m["madde_no_2019"], _m["madde_no_2023"]
    if _no19 is not None and _no23 is not None and _no19 != _no23:
        _EXPECTED_STRUCTURAL_COUNTS["RENUMBERED"] += 1
    # NEDEN MOVED her zaman 0: ground_truth.json'da gerçek bir MOVED örneği
    # yok (bkz. docs/ARCHITECTURE.md §7 -- bilinen sınırlama).
_EXPECTED_TOTAL = len(_GROUND_TRUTH["maddeler"])
_DETAIL_COLUMNS = [
    "Değişim Türü",
    "Numarası Değişti",
    "Yeri Değişti",
    "Eski Madde No",
    "Yeni Madde No",
    "Başlık",
    "Eski Bölüm",
    "Yeni Bölüm",
    "Eski Metin",
    "Yeni Metin",
]

_OLD_BYTES = (SAMPLES_DIR / "yonetmelik_2019.pdf").read_bytes()
_NEW_BYTES = (SAMPLES_DIR / "yonetmelik_2023.pdf").read_bytes()
_RESULT = compare_documents(_OLD_BYTES, "yonetmelik_2019.pdf", _NEW_BYTES, "yonetmelik_2023.pdf")


class TestMetricsSummaryDataFrame:
    def test_her_change_type_icin_bir_satir(self):
        df = build_summary_dataframe(_RESULT)
        assert len(df) == len(ChangeType)
        assert list(df.columns) == ["Değişim Türü", "Sayı", "Yüzde"]

    def test_sayilar_ground_truth_ile_esler(self):
        df = build_summary_dataframe(_RESULT)
        by_label = dict(zip(df["Değişim Türü"], df["Sayı"]))
        for change_type, expected in _EXPECTED_COUNTS.items():
            label = CHANGE_TYPE_LABELS_TR[change_type.value]
            assert by_label[label] == expected, f"{label}: beklenen {expected}, bulunan {by_label[label]}"

    def test_yuzdeler_toplami_yuz_civarinda(self):
        # NEDEN dar olmayan bir tolerans: her yüzde ayrı ayrı 1 ondalığa
        # YUVARLANDIĞI için toplam, kayan nokta hassasiyeti VE yuvarlama payı
        # yüzünden 100'den birkaç onda bir sapabilir -- tek tek yüzdelerin
        # DOĞRULUĞU zaten test_yuzdeler_dogru_hesaplanir'de (summary_builder
        # testleri) ayrı ayrı kontrol ediliyor.
        df = build_summary_dataframe(_RESULT)
        assert 99.5 <= df["Yüzde"].sum() <= 100.5

    def test_bos_sonucta_bile_dort_satir_ve_dogru_kolonlar(self):
        from src.analysis.pipeline import ComparisonResult
        from src.models import DocumentMeta

        empty = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[],
        )
        df = build_summary_dataframe(empty)
        assert len(df) == 4  # IDENTICAL/MODIFIED/ADDED/REMOVED
        assert list(df.columns) == ["Değişim Türü", "Sayı", "Yüzde"]
        assert (df["Sayı"] == 0).all()
        assert (df["Yüzde"] == 0.0).all()


class TestMetricsStructuralDataFrame:
    """
    build_structural_dataframe -- İÇERİK durumundan BAĞIMSIZ, "Numarası
    Değişti"/"Yeri Değişti" sayımı (bkz. classifier.py NEDEN notu).
    """

    def test_iki_satir_dogru_kolonlar(self):
        df = build_structural_dataframe(_RESULT)
        assert len(df) == 2
        assert list(df.columns) == ["Yapısal Değişiklik", "Sayı"]

    def test_sayilar_ground_truth_ile_esler(self):
        # NEDEN 5 (4 DEĞİL): eski modelde SADECE saf RENUMBERED maddeler
        # (4 tanesi) sayılırdı -- birim_sorumluluklari (9->8) hem numarası
        # HEM içeriği değiştiği için MODIFIED'a düşüp bu sayımdan
        # KAYBOLURDU. Yeni modelde numarasi_degisti içerikten BAĞIMSIZ
        # hesaplandığı için o madde de burada sayılır (bkz. Adım 9'un TAM
        # konusu, _EXPECTED_STRUCTURAL_COUNTS NEDEN notu).
        df = build_structural_dataframe(_RESULT)
        by_label = dict(zip(df["Yapısal Değişiklik"], df["Sayı"]))
        assert by_label[CHANGE_TYPE_LABELS_TR["RENUMBERED"]] == _EXPECTED_STRUCTURAL_COUNTS["RENUMBERED"] == 5
        assert by_label[CHANGE_TYPE_LABELS_TR["MOVED"]] == _EXPECTED_STRUCTURAL_COUNTS["MOVED"] == 0

    def test_bos_sonucta_bile_iki_satir_sifir_sayi(self):
        from src.analysis.pipeline import ComparisonResult
        from src.models import DocumentMeta

        empty = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[],
        )
        df = build_structural_dataframe(empty)
        assert len(df) == 2
        assert (df["Sayı"] == 0).all()


class TestMetricsSectionBreakdown:
    """
    build_section_breakdown -- BÖLÜM bazında İÇERİK DURUMU dökümü (bkz.
    metrics.py NEDEN notu). ground_truth.json'daki 2019/2023 örneğinde
    dört bölüm var: BİRİNCİ/İKİNCİ/ÜÇÜNCÜ/DÖRDÜNCÜ BÖLÜM.
    """

    def test_kolonlar_dogru(self):
        df = build_section_breakdown(_RESULT)
        assert list(df.columns) == ["Bölüm", "Değişmedi", "Değişti", "Yeni Eklendi", "Kaldırıldı"]

    def test_bolum_sirasi_belge_sirasidir(self):
        df = build_section_breakdown(_RESULT)
        assert list(df["Bölüm"]) == ["BİRİNCİ BÖLÜM", "İKİNCİ BÖLÜM", "ÜÇÜNCÜ BÖLÜM", "DÖRDÜNCÜ BÖLÜM"]

    def test_sayilar_ground_truth_ile_esler(self):
        # NEDEN bu tam sayılar: gerçek 2019/2023 örneğinde her bölümün
        # içeriği BİLİNEN bir dağılıma sahip (bkz. sınıf docstring'i);
        # BİRİNCİ BÖLÜM (3 Değişmedi + 1 Değişti) kullanıcının kendi
        # verdiği örnekle BİREBİR örtüşüyor.
        df = build_section_breakdown(_RESULT)
        by_bolum = df.set_index("Bölüm")
        assert by_bolum.loc["BİRİNCİ BÖLÜM"].to_dict() == {
            "Değişmedi": 3, "Değişti": 1, "Yeni Eklendi": 0, "Kaldırıldı": 0,
        }
        assert by_bolum.loc["İKİNCİ BÖLÜM"].to_dict() == {
            "Değişmedi": 1, "Değişti": 2, "Yeni Eklendi": 0, "Kaldırıldı": 1,
        }
        assert by_bolum.loc["ÜÇÜNCÜ BÖLÜM"].to_dict() == {
            "Değişmedi": 1, "Değişti": 1, "Yeni Eklendi": 2, "Kaldırıldı": 0,
        }
        assert by_bolum.loc["DÖRDÜNCÜ BÖLÜM"].to_dict() == {
            "Değişmedi": 3, "Değişti": 0, "Yeni Eklendi": 0, "Kaldırıldı": 0,
        }

    def test_toplam_result_counts_ile_esler(self):
        # NEDEN kritik: hiçbir maddenin bölüm dökümünde SESSİZCE
        # kaybolmadığını (bkz. BOLUMSUZ_ETIKETI NEDEN notu) doğrular --
        # tablodaki TÜM hücrelerin toplamı result.counts() toplamına
        # EŞİT olmalı.
        df = build_section_breakdown(_RESULT)
        toplam_tablo = df[["Değişmedi", "Değişti", "Yeni Eklendi", "Kaldırıldı"]].to_numpy().sum()
        assert toplam_tablo == sum(_RESULT.counts().values()) == _EXPECTED_TOTAL

    def test_bolum_bilgisi_olmayan_madde_ayri_kovada_toplanir(self):
        from src.analysis.classifier import ClassifiedSection
        from src.analysis.pipeline import ComparisonResult, ComparisonRow
        from src.models import DocumentMeta, Section, TextUnit

        unit = TextUnit(doc_id="t", page_no=1, block_index=0, char_start=0, char_end=1, heading_path=(), text="x")
        # NEDEN heading_path tek elemanlı ("MADDE 1"): [:-1] alındığında
        # KISIM/BÖLÜM zinciri BOŞ kalır -- bölüm bilgisi olmayan madde budur.
        section = Section(
            doc_id="t", heading_path=("MADDE 1",), section_type="MADDE",
            madde_no=1, baslik="X", order_index=0, units=[unit],
        )
        result = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[ComparisonRow(classified=ClassifiedSection(change_type=ChangeType.IDENTICAL, old=section, new=section), section_diff=None)],
        )
        df = build_section_breakdown(result)
        assert list(df["Bölüm"]) == [BOLUMSUZ_ETIKETI]
        assert df.iloc[0]["Değişmedi"] == 1

    def test_bos_sonucta_bos_dataframe_dogru_kolonlarla(self):
        from src.analysis.pipeline import ComparisonResult
        from src.models import DocumentMeta

        empty = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[],
        )
        df = build_section_breakdown(empty)
        assert len(df) == 0
        assert list(df.columns) == ["Bölüm", "Değişmedi", "Değişti", "Yeni Eklendi", "Kaldırıldı"]


class TestMetricsDetailDataFrame:
    def test_satir_sayisi_ground_truth_ile_esler(self):
        df = build_detail_dataframe(_RESULT)
        assert len(df) == _EXPECTED_TOTAL

    def test_kolonlar_beklenen_sirada(self):
        df = build_detail_dataframe(_RESULT)
        assert list(df.columns) == _DETAIL_COLUMNS

    def test_removed_satirda_yeni_madde_no_bos(self):
        df = build_detail_dataframe(_RESULT)
        removed = df[df["Değişim Türü"] == CHANGE_TYPE_LABELS_TR["REMOVED"]]
        assert len(removed) == 1
        assert removed.iloc[0]["Yeni Madde No"] is None or pd.isna(removed.iloc[0]["Yeni Madde No"])
        assert removed.iloc[0]["Eski Madde No"] == 8

    def test_added_satirda_eski_madde_no_bos(self):
        df = build_detail_dataframe(_RESULT)
        added = df[df["Değişim Türü"] == CHANGE_TYPE_LABELS_TR["ADDED"]]
        assert len(added) == 2
        assert added["Eski Madde No"].isna().all()

    def test_birim_sorumluluklari_satirinda_numarasi_degisti_evet(self):
        # NEDEN kritik: Adım 9'un TAM konusu -- bu madde hem İÇERİK hem
        # NUMARA olarak değişti (bkz. ground_truth.json/birim_sorumlulukları,
        # 2019 no.9 -> 2023 no.8); detay tablosunda İKİSİ de görünmeli.
        df = build_detail_dataframe(_RESULT)
        satir = df[df["Başlık"] == "Birim Sorumlulukları"].iloc[0]
        assert satir["Değişim Türü"] == CHANGE_TYPE_LABELS_TR["MODIFIED"]
        assert satir["Numarası Değişti"] == "Evet"
        assert satir["Yeri Değişti"] == "Hayır"

    def test_yururluk_satirinda_icerik_ayni_ama_numarasi_degisti_evet(self):
        df = build_detail_dataframe(_RESULT)
        satir = df[df["Başlık"] == "Yürürlük"].iloc[0]
        assert satir["Değişim Türü"] == CHANGE_TYPE_LABELS_TR["IDENTICAL"]
        assert satir["Numarası Değişti"] == "Evet"

    def test_amac_satirinda_hicbir_sey_degismedi(self):
        df = build_detail_dataframe(_RESULT)
        satir = df[df["Başlık"] == "Amaç"].iloc[0]
        assert satir["Değişim Türü"] == CHANGE_TYPE_LABELS_TR["IDENTICAL"]
        assert satir["Numarası Değişti"] == "Hayır"
        assert satir["Yeri Değişti"] == "Hayır"

    def test_bos_sonucta_bile_kolon_basliklari_korunur(self):
        from src.analysis.pipeline import ComparisonResult
        from src.models import DocumentMeta

        empty = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[],
        )
        df = build_detail_dataframe(empty)
        assert len(df) == 0
        assert list(df.columns) == _DETAIL_COLUMNS


def _tek_maddelik_sonuc(metin: str, *, baslik: str = "Test") -> "ComparisonResult":
    """Tek bir ADDED Section içeren minimal ComparisonResult -- güvenlik testleri için."""
    from src.analysis.classifier import ClassifiedSection
    from src.analysis.pipeline import ComparisonResult, ComparisonRow
    from src.models import DocumentMeta, Section, TextUnit

    unit = TextUnit(
        doc_id="t", page_no=1, block_index=0, char_start=0, char_end=len(metin), heading_path=(), text=metin
    )
    section = Section(
        doc_id="t", heading_path=("MADDE 1",), section_type="MADDE",
        madde_no=1, baslik=baslik, order_index=0, units=[unit],
    )
    return ComparisonResult(
        old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
        new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
        rows=[ComparisonRow(classified=ClassifiedSection(change_type=ChangeType.ADDED, old=None, new=section), section_diff=None)],
    )


class TestMetricsFormulaEnjeksiyonuKorumasi:
    """
    Güvenlik regresyon testleri (CWE-1236 / "CSV Injection"): belge
    içeriğinden gelen bir madde metni "=", "+", "-" ya da "@" ile
    BAŞLIYORSA (örn. bir liste öğesi "- İlgili birimler..." gibi doğal
    biçimde bile olabilir), dışa aktarılan CSV/Excel'de bu bir FORMÜL
    olarak yorumlanmamalı -- Excel/LibreOffice bu karakterleri hücre
    başında formül tetikleyicisi sayar.
    """

    @pytest.mark.parametrize("tetikleyici", ["=", "+", "-", "@"])
    def test_tehlikeli_karakterle_baslayan_metin_tek_tirnakla_kacirilir(self, tetikleyici):
        tehlikeli = f"{tetikleyici}cmd|'/c calc'!A1"
        result = _tek_maddelik_sonuc(tehlikeli)

        df = build_detail_dataframe(result)

        assert df.iloc[0]["Yeni Metin"] == f"'{tehlikeli}"

    def test_zararsiz_metin_degistirilmeden_kalir(self):
        result = _tek_maddelik_sonuc("(1) Kurum verileri paylaşılamaz.")

        df = build_detail_dataframe(result)

        assert df.iloc[0]["Yeni Metin"] == "(1) Kurum verileri paylaşılamaz."

    def test_metin_ortasindaki_tehlikeli_karakter_dokunulmaz(self):
        # NEDEN: formül yorumlanması SADECE hücrenin İLK karakteriyle
        # ilgilidir -- metin içinde herhangi bir yerde "=" geçmesi risk
        # OLUŞTURMAZ, gereksiz kaçırma OKUNABİLİRLİĞİ bozardı.
        result = _tek_maddelik_sonuc("Bu madde x=5 formülünü içerir.")

        df = build_detail_dataframe(result)

        assert df.iloc[0]["Yeni Metin"] == "Bu madde x=5 formülünü içerir."

    def test_baslik_ve_bolum_de_korunur(self):
        result = _tek_maddelik_sonuc("normal metin", baslik="=HYPERLINK(\"http://kotu\")")

        df = build_detail_dataframe(result)

        assert df.iloc[0]["Başlık"] == "'=HYPERLINK(\"http://kotu\")"

    def test_csv_disa_aktarimda_da_kacirilmis_halde_gorunur(self):
        result = _tek_maddelik_sonuc("-Ilgili birimler bildirmekle yukumludur.")

        csv_text = export_to_csv(result).decode("utf-8-sig")

        assert "'-Ilgili birimler" in csv_text

    def test_excel_disa_aktarimda_da_kacirilmis_halde_gorunur(self):
        import openpyxl

        result = _tek_maddelik_sonuc("=cmd|'/c calc'!A1")

        xlsx_bytes = export_to_excel(result)
        wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
        detay = wb["Detay"]
        header = [c.value for c in next(detay.iter_rows(min_row=1, max_row=1))]
        row = [c.value for c in next(detay.iter_rows(min_row=2, max_row=2))]
        deger = dict(zip(header, row))

        assert deger["Yeni Metin"] == "'=cmd|'/c calc'!A1"


class TestSummaryBuilderStats:
    def test_toplamlar_ground_truth_ile_esler(self):
        stats = build_summary_stats(_RESULT)
        assert stats.toplam_karsilastirma == _EXPECTED_TOTAL
        assert stats.eski_toplam_madde == 13  # 2019: 13 madde
        assert stats.yeni_toplam_madde == 14  # 2023: 14 madde

    def test_sayilar_ground_truth_ile_esler(self):
        stats = build_summary_stats(_RESULT)
        assert stats.sayilar == _EXPECTED_COUNTS

    def test_yuzdeler_dogru_hesaplanir(self):
        stats = build_summary_stats(_RESULT)
        for change_type, count in _EXPECTED_COUNTS.items():
            expected_pct = round(count / _EXPECTED_TOTAL * 100, 1)
            assert stats.yuzdeler[change_type] == expected_pct

    def test_en_cok_degisen_bolum_ucuncu_bolumdur(self):
        # NEDEN ÜÇÜNCÜ BÖLÜM: birim_sorumluluklari (MODIFIED içerik +
        # numarası değişti), ust_yonetim_sorumlulugu (İÇERİK IDENTICAL ama
        # numarası değişti -- yine de "bir şey değişti" sayılır, bkz.
        # summary_builder.py::_en_cok_degisen_bolum NEDEN notu),
        # veri_yonetisim_kurulu(ADDED), acik_veri_portali(ADDED) -- 4 madde,
        # ground truth'taki HERHANGİ başka bir bölümden fazla.
        stats = build_summary_stats(_RESULT)
        assert stats.en_cok_degisen_bolumler == [("ÜÇÜNCÜ BÖLÜM", 4)]

    def test_yapisal_sayilar_ground_truth_ile_esler(self):
        stats = build_summary_stats(_RESULT)
        assert stats.yapisal_sayilar == _EXPECTED_STRUCTURAL_COUNTS

    def test_hicbir_degisiklik_yoksa_en_cok_degisen_bolum_nonedir(self):
        from src.analysis.classifier import ClassifiedSection
        from src.analysis.pipeline import ComparisonResult, ComparisonRow
        from src.models import DocumentMeta, Section, TextUnit

        unit = TextUnit(doc_id="t", page_no=1, block_index=0, char_start=0, char_end=1, heading_path=(), text="x")
        section = Section(
            doc_id="t", heading_path=("BİRİNCİ BÖLÜM", "MADDE 1"), section_type="MADDE",
            madde_no=1, baslik="X", order_index=0, units=[unit],
        )
        result = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[ComparisonRow(classified=ClassifiedSection(change_type=ChangeType.IDENTICAL, old=section, new=section), section_diff=None)],
        )
        stats = build_summary_stats(result)
        assert stats.en_cok_degisen_bolumler == []

    def test_birden_fazla_bolum_esitse_hepsi_listelenir(self):
        # NEDEN: "en çok değişen bölüm" TEK bir kazanana rastgele/dict-
        # sıralamasına göre indirgenmemeli -- iki bölüm AYNI sayıda
        # değişiklik taşıyorsa İKİSİ de listelenmeli (bkz.
        # summary_builder.py::_en_cok_degisen_bolumler NEDEN notu).
        from src.analysis.classifier import ClassifiedSection
        from src.analysis.pipeline import ComparisonResult, ComparisonRow
        from src.models import DocumentMeta, Section, TextUnit

        def _section(heading_path, madde_no):
            unit = TextUnit(doc_id="t", page_no=1, block_index=0, char_start=0, char_end=1, heading_path=(), text="x")
            return Section(
                doc_id="t", heading_path=heading_path, section_type="MADDE",
                madde_no=madde_no, baslik="X", order_index=madde_no, units=[unit],
            )

        a_old = _section(("A BÖLÜMÜ", "MADDE 1"), 1)
        a_new = _section(("A BÖLÜMÜ", "MADDE 1"), 2)  # numarası değişti
        b_old = _section(("B BÖLÜMÜ", "MADDE 2"), 1)
        b_new = _section(("B BÖLÜMÜ", "MADDE 2"), 2)  # numarası değişti

        result = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[
                ComparisonRow(
                    classified=ClassifiedSection(
                        change_type=ChangeType.IDENTICAL, old=a_old, new=a_new, numarasi_degisti=True
                    ),
                    section_diff=None,
                ),
                ComparisonRow(
                    classified=ClassifiedSection(
                        change_type=ChangeType.IDENTICAL, old=b_old, new=b_new, numarasi_degisti=True
                    ),
                    section_diff=None,
                ),
            ],
        )
        stats = build_summary_stats(result)
        assert stats.en_cok_degisen_bolumler == [("A BÖLÜMÜ", 1), ("B BÖLÜMÜ", 1)]

    def test_sonuc_tipi(self):
        assert isinstance(build_summary_stats(_RESULT), SummaryStats)


class TestSummaryBuilderOneCikanDegisiklikler:
    """
    "Öne Çıkan Değişiklikler" -- SADECE metinsel olarak gözlemlenebilen
    olgulara dayalı, şablon tabanlı kısa cümle listesi (bkz.
    summary_builder.py::_one_cikan_degisiklikler NEDEN notu).
    """

    def test_gercek_veri_setinde_beklenen_cumleler_gecer(self):
        # NEDEN bu ÜÇ cümle: ground_truth.json'da tam olarak bu maddeler
        # ilgili kategoriye düşer (bkz. TestMetricsDetailDataFrame ve
        # kullanıcı geri bildiriminin ÖRNEK cümleleriyle birebir örtüşür).
        stats = build_summary_stats(_RESULT)
        assert "Veri Paylaşımı maddesinde içerik değişikliği tespit edildi." in stats.one_cikan_degisiklikler
        assert "Arşivleme Esasları maddesi yeni belgede bulunmuyor." in stats.one_cikan_degisiklikler
        assert "Veri Yönetişim Kurulu maddesi yeni belgede eklendi." in stats.one_cikan_degisiklikler
        assert f"{_EXPECTED_STRUCTURAL_COUNTS['RENUMBERED']} maddenin numarası değişti." in stats.one_cikan_degisiklikler

    def test_hicbir_yorum_kelimesi_gecmez(self):
        # NEDEN kritik: kullanıcı geri bildiriminin AÇIKÇA yasakladığı
        # hukuki/normatif yorum kelimeleri hiçbir şablonda YER ALMAMALI.
        stats = build_summary_stats(_RESULT)
        yasakli_kelimeler = [
            "zayıflat", "güçlendir", "risk", "hukuken", "önemli", "kapsamlı", "olumsuz", "olumlu",
        ]
        for cumle in stats.one_cikan_degisiklikler:
            for kelime in yasakli_kelimeler:
                assert kelime not in cumle.lower(), f"yasaklı kelime '{kelime}' şu cümlede bulundu: {cumle!r}"

    def test_moved_sifirsa_yeri_degisti_cumlesi_yok(self):
        stats = build_summary_stats(_RESULT)
        assert not any("yeri değişti" in c for c in stats.one_cikan_degisiklikler)

    def test_hicbir_degisiklik_yoksa_bos_liste(self):
        from src.analysis.classifier import ClassifiedSection
        from src.analysis.pipeline import ComparisonResult, ComparisonRow
        from src.models import DocumentMeta, Section, TextUnit

        unit = TextUnit(doc_id="t", page_no=1, block_index=0, char_start=0, char_end=1, heading_path=(), text="x")
        section = Section(
            doc_id="t", heading_path=("BİRİNCİ BÖLÜM", "MADDE 1"), section_type="MADDE",
            madde_no=1, baslik="X", order_index=0, units=[unit],
        )
        result = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=[ComparisonRow(classified=ClassifiedSection(change_type=ChangeType.IDENTICAL, old=section, new=section), section_diff=None)],
        )
        stats = build_summary_stats(result)
        assert stats.one_cikan_degisiklikler == []

    def test_kategori_basina_ust_sinir_ve_toplu_ek_cumle(self):
        from src.analysis.classifier import ClassifiedSection
        from src.analysis.pipeline import ComparisonResult, ComparisonRow
        from src.models import DocumentMeta, Section, TextUnit

        rows = []
        for i in range(7):  # SUMMARY_HIGHLIGHT_MAX_ITEMS_PER_CATEGORY (5) + 2
            unit = TextUnit(doc_id="t", page_no=1, block_index=0, char_start=0, char_end=1, heading_path=(), text="x")
            section = Section(
                doc_id="t", heading_path=("BÖLÜM", f"MADDE {i}"), section_type="MADDE",
                madde_no=i, baslik=f"Madde{i}", order_index=i, units=[unit],
            )
            rows.append(
                ComparisonRow(classified=ClassifiedSection(change_type=ChangeType.ADDED, old=None, new=section), section_diff=None)
            )
        result = ComparisonResult(
            old_meta=DocumentMeta(doc_id="x", source_path="x", file_type="pdf"),
            new_meta=DocumentMeta(doc_id="y", source_path="y", file_type="pdf"),
            rows=rows,
        )
        stats = build_summary_stats(result)
        eklenen_cumleler = [c for c in stats.one_cikan_degisiklikler if "eklen" in c]
        assert len(eklenen_cumleler) == 6  # 5 tek tek + 1 toplu
        assert "Yeni belgede eklenen 2 madde daha var." in stats.one_cikan_degisiklikler

    def test_deterministik_iki_cagri_ayni_sonucu_verir(self):
        assert build_summary_stats(_RESULT).one_cikan_degisiklikler == build_summary_stats(_RESULT).one_cikan_degisiklikler


class TestSummaryBuilderMarkdown:
    def test_tum_icerik_ve_yapisal_etiketler_gecer(self):
        stats = build_summary_stats(_RESULT)
        md = render_summary_markdown(stats)
        for change_type in ChangeType:
            assert CHANGE_TYPE_LABELS_TR[change_type.value] in md
        assert CHANGE_TYPE_LABELS_TR["RENUMBERED"] in md
        assert CHANGE_TYPE_LABELS_TR["MOVED"] in md

    def test_icerik_ve_yapisal_iki_ayri_baslik_altinda(self):
        # NEDEN: bu ikisi BAĞIMSIZ boyutlardır, TEK bir listede
        # birleştirilmemeli (bkz. modül NEDEN notu) -- "İçerik Durumu" ve
        # "Yapısal Değişiklikler" AYRI başlıklar altında görünmeli.
        stats = build_summary_stats(_RESULT)
        md = render_summary_markdown(stats)
        assert "### İçerik Durumu" in md
        assert "### Yapısal Değişiklikler" in md
        assert md.index("### İçerik Durumu") < md.index("### Yapısal Değişiklikler")

    def test_yapisal_cumleler_istenen_ifadeyle_eslesir(self):
        # NEDEN: kullanıcının BİREBİR istediği ifade -- "X maddede madde
        # numarası değişikliği tespit edildi."
        stats = build_summary_stats(_RESULT)
        md = render_summary_markdown(stats)
        assert f"{_EXPECTED_STRUCTURAL_COUNTS['RENUMBERED']} maddede madde numarası değişikliği tespit edildi." in md
        assert f"{_EXPECTED_STRUCTURAL_COUNTS['MOVED']} maddede yeri değişikliği tespit edildi." in md

    def test_sayilar_metinde_gecer(self):
        stats = build_summary_stats(_RESULT)
        md = render_summary_markdown(stats)
        assert "**13**" in md
        assert "**14**" in md
        assert "**15**" in md

    def test_en_cok_degisen_bolum_metinde_gecer(self):
        stats = build_summary_stats(_RESULT)
        md = render_summary_markdown(stats)
        assert "ÜÇÜNCÜ BÖLÜM" in md

    def test_bolum_yoksa_o_satir_hic_gorunmez(self):
        stats = SummaryStats(
            toplam_karsilastirma=0,
            eski_toplam_madde=0,
            yeni_toplam_madde=0,
            sayilar=dict.fromkeys(ChangeType, 0),
            yuzdeler=dict.fromkeys(ChangeType, 0.0),
            yapisal_sayilar={"RENUMBERED": 0, "MOVED": 0},
            en_cok_degisen_bolumler=[],
        )
        md = render_summary_markdown(stats)
        assert "En çok değişiklik" not in md

    def test_birden_fazla_bolum_esitse_ikisi_de_ve_cogul_ek_gorunur(self):
        stats = SummaryStats(
            toplam_karsilastirma=2,
            eski_toplam_madde=2,
            yeni_toplam_madde=2,
            sayilar=dict.fromkeys(ChangeType, 0),
            yuzdeler=dict.fromkeys(ChangeType, 0.0),
            yapisal_sayilar={"RENUMBERED": 0, "MOVED": 0},
            en_cok_degisen_bolumler=[("A BÖLÜMÜ", 2), ("B BÖLÜMÜ", 2)],
        )
        md = render_summary_markdown(stats)
        assert "A BÖLÜMÜ" in md and "B BÖLÜMÜ" in md
        assert "bölümlerinde görüldü (2 madde)" in md

    def test_one_cikan_degisiklikler_bolumu_bos_ise_gorunmez(self):
        stats = SummaryStats(
            toplam_karsilastirma=0,
            eski_toplam_madde=0,
            yeni_toplam_madde=0,
            sayilar=dict.fromkeys(ChangeType, 0),
            yuzdeler=dict.fromkeys(ChangeType, 0.0),
            yapisal_sayilar={"RENUMBERED": 0, "MOVED": 0},
            one_cikan_degisiklikler=[],
        )
        md = render_summary_markdown(stats)
        assert "Öne Çıkan Değişiklikler" not in md

    def test_one_cikan_degisiklikler_gercek_veride_gorunur(self):
        stats = build_summary_stats(_RESULT)
        md = render_summary_markdown(stats)
        assert "### Öne Çıkan Değişiklikler" in md
        assert "- Veri Paylaşımı maddesinde içerik değişikliği tespit edildi." in md

    def test_deterministik_iki_cagri_ayni_sonucu_verir(self):
        stats = build_summary_stats(_RESULT)
        assert render_summary_markdown(stats) == render_summary_markdown(stats)


class TestSummaryHtml:
    def test_ozel_karakterli_bolum_adi_kacirilir(self):
        stats = SummaryStats(
            toplam_karsilastirma=1,
            eski_toplam_madde=1,
            yeni_toplam_madde=1,
            sayilar=dict.fromkeys(ChangeType, 0),
            yuzdeler=dict.fromkeys(ChangeType, 0.0),
            yapisal_sayilar={"RENUMBERED": 0, "MOVED": 0},
            en_cok_degisen_bolumler=[("<script>alert(1)</script>", 3)],
        )
        html = render_summary_html(stats)
        assert "<script>alert(1)</script>" not in html
        assert "&lt;script&gt;" in html

    def test_one_cikan_degisiklikler_ozel_karakter_kacirilir(self):
        stats = SummaryStats(
            toplam_karsilastirma=1,
            eski_toplam_madde=1,
            yeni_toplam_madde=1,
            sayilar=dict.fromkeys(ChangeType, 0),
            yuzdeler=dict.fromkeys(ChangeType, 0.0),
            yapisal_sayilar={"RENUMBERED": 0, "MOVED": 0},
            one_cikan_degisiklikler=['<script>alert(1)</script> maddesi yeni belgede eklendi.'],
        )
        html = render_summary_html(stats)
        assert "<script>alert(1)</script>" not in html
        assert "&lt;script&gt;" in html

    def test_one_cikan_degisiklikler_gercek_veride_gorunur(self):
        stats = build_summary_stats(_RESULT)
        html = render_summary_html(stats)
        assert "<h3>Öne Çıkan Değişiklikler</h3>" in html
        assert "<li>Arşivleme Esasları maddesi yeni belgede bulunmuyor.</li>" in html

    def test_her_icerik_ve_yapisal_rozet_gecer(self):
        stats = build_summary_stats(_RESULT)
        html = render_summary_html(stats)
        for change_type in ChangeType:
            assert f"mk-badge-{change_type.value.lower()}" in html
        assert "mk-badge-renumbered" in html
        assert "mk-badge-moved" in html

    def test_yapisal_sayilar_html_metninde_gecer(self):
        stats = build_summary_stats(_RESULT)
        html = render_summary_html(stats)
        assert f"{_EXPECTED_STRUCTURAL_COUNTS['RENUMBERED']} maddede madde numarası değişikliği tespit edildi." in html


class TestExporterExcel:
    def test_uc_sayfa_uretir_dogru_boyutlarda(self):
        # NEDEN üç (iki DEĞİL): "Özet" (İÇERİK durumu) + "Yapısal" (Numarası/
        # Yeri Değişti -- bkz. classifier.py NEDEN notu) + "Detay".
        xlsx_bytes = export_to_excel(_RESULT)
        sheets = pd.read_excel(BytesIO(xlsx_bytes), sheet_name=None)
        assert set(sheets.keys()) == {"Özet", "Yapısal", "Detay"}
        assert sheets["Özet"].shape == (len(ChangeType), 3)
        assert sheets["Yapısal"].shape == (2, 2)
        assert sheets["Detay"].shape[0] == _EXPECTED_TOTAL

    def test_detay_sayfasi_sayimla_tutarli(self):
        xlsx_bytes = export_to_excel(_RESULT)
        detay = pd.read_excel(BytesIO(xlsx_bytes), sheet_name="Detay")
        for change_type, expected in _EXPECTED_COUNTS.items():
            label = CHANGE_TYPE_LABELS_TR[change_type.value]
            assert (detay["Değişim Türü"] == label).sum() == expected

    def test_yapisal_sayfa_sayimla_tutarli(self):
        xlsx_bytes = export_to_excel(_RESULT)
        yapisal = pd.read_excel(BytesIO(xlsx_bytes), sheet_name="Yapısal")
        by_label = dict(zip(yapisal["Yapısal Değişiklik"], yapisal["Sayı"]))
        assert by_label[CHANGE_TYPE_LABELS_TR["RENUMBERED"]] == _EXPECTED_STRUCTURAL_COUNTS["RENUMBERED"]
        assert by_label[CHANGE_TYPE_LABELS_TR["MOVED"]] == _EXPECTED_STRUCTURAL_COUNTS["MOVED"]


class TestExporterCsv:
    def test_utf8_sig_bom_ile_baslar(self):
        csv_bytes = export_to_csv(_RESULT)
        assert csv_bytes.startswith(b"\xef\xbb\xbf")

    def test_baslik_satiri_dogru(self):
        csv_bytes = export_to_csv(_RESULT)
        first_line = csv_bytes.decode("utf-8-sig").splitlines()[0]
        assert first_line == ",".join(_DETAIL_COLUMNS)

    def test_satir_sayisi_basliklar_dahil(self):
        csv_bytes = export_to_csv(_RESULT)
        lines = csv_bytes.decode("utf-8-sig").splitlines()
        assert len(lines) == _EXPECTED_TOTAL + 1  # +1 başlık satırı

    def test_geri_okunabilir_pandas_ile(self):
        csv_bytes = export_to_csv(_RESULT)
        df = pd.read_csv(BytesIO(csv_bytes), encoding="utf-8-sig")
        assert len(df) == _EXPECTED_TOTAL


class TestExporterHtml:
    def test_gecerli_bagimsiz_html_belgesi(self):
        html_bytes = export_to_html(_RESULT)
        html_text = html_bytes.decode("utf-8")
        assert html_text.startswith("<!doctype html>")
        assert "<style>" in html_text and "</style>" in html_text
        assert "mk-badge-modified" in html_text  # CSS enjekte edilmiş

    def test_css_tum_icerik_ve_yapisal_renklerini_icerir(self):
        html_text = export_to_html(_RESULT).decode("utf-8")
        for change_type in ChangeType:
            assert f".mk-badge-{change_type.value.lower()}" in html_text
        assert ".mk-badge-renumbered" in html_text
        assert ".mk-badge-moved" in html_text

    def test_her_maddenin_basligi_gecer(self):
        html_text = export_to_html(_RESULT).decode("utf-8")
        for m in _GROUND_TRUTH["maddeler"]:
            assert m["baslik"] in html_text

    def test_ozet_bolumu_icerir(self):
        html_text = export_to_html(_RESULT).decode("utf-8")
        assert "Yönetici Özeti" in html_text

    def test_kaldirilan_madde_kirmizi_sinifla_gorunur(self):
        html_text = export_to_html(_RESULT).decode("utf-8")
        assert "mk-diff-removed" in html_text

    def test_yapisal_rozetler_satir_basliklarinda_da_gorunur(self):
        # NEDEN kritik: export_to_html, render_row_header_html'i
        # numarasi_degisti/yeri_degisti PARAMETRELERİ OLMADAN çağırıyordu
        # (bu Adım 9 sırasında fark edilip düzeltilen bir hataydı) --
        # yapısal rozetler SADECE Yönetici Özeti'nde değil, HER maddenin
        # kendi başlığında da görünmeli (bkz. html_renderer.py::
        # render_row_header_html).
        html_text = export_to_html(_RESULT).decode("utf-8")
        # "Yürürlük" satırı: içerik IDENTICAL + numarası değişti.
        yururluk_index = html_text.index("Yürürlük<")
        civar = html_text[yururluk_index : yururluk_index + 400]
        assert "mk-badge-renumbered" in civar, "Yürürlük satırının başlığında 'numarası değişti' rozeti yok"

    def test_html_ozel_karakter_icermez_kacirilmamis(self):
        # NEDEN: gövde metninde "<" "&" gibi karakterler varsa (bkz.
        # dayanak maddesi "T.C." kısaltması İÇERMEZ ama genel ilke olarak)
        # ham HTML olarak sızmamalı -- html_renderer zaten escape ediyor,
        # burada dışa aktarılan dosyanın GERÇEKTEN html.escape() çıktısı
        # taşıdığını (örn. hiç kapanmamış bir <div> kalmadığını) doğruluyoruz.
        html_text = export_to_html(_RESULT).decode("utf-8")
        assert html_text.count("<div") == html_text.count("</div>")

    @pytest.mark.parametrize("dosya_2019,dosya_2023", [("yonetmelik_2019.docx", "yonetmelik_2023.docx")])
    def test_docx_ile_de_calisir(self, dosya_2019, dosya_2023):
        old_bytes = (SAMPLES_DIR / dosya_2019).read_bytes()
        new_bytes = (SAMPLES_DIR / dosya_2023).read_bytes()
        result = compare_documents(old_bytes, dosya_2019, new_bytes, dosya_2023)
        html_text = export_to_html(result).decode("utf-8")
        assert html_text.startswith("<!doctype html>")
        assert "Yönetici Özeti" in html_text
