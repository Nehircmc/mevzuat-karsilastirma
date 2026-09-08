"""
src/analysis/pipeline.py::compare_documents testleri -- özellikle
classify_content() (İÇERİK durumu, kaba karakter/kelime benzerliği) ile
diff_section_match() (ayrıntılı cümle/kelime düzeyi) arasındaki ETKİLEŞİMİ
kapsayan, tek başlarına classifier.py/differ.py testleriyle YAKALANAMAYAN
regresyonlar için.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import docx as docx_lib

from src.analysis.pipeline import compare_documents
from src.models import ChangeType


def _build_madde_docx_bytes(text: str) -> bytes:
    """TEK bir MADDE 1 içeren minimal bir DOCX üretir (verilen gövde metniyle)."""
    document = docx_lib.Document()
    document.add_paragraph(f"MADDE 1- {text}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_docx_bytes(*, degistir: bool) -> bytes:
    """
    TEK bir çok uzun (~60 fıkralı) madde içeren minimal bir DOCX üretir --
    `degistir=True` ise fıkra 30'daki TEK bir ifade değiştirilir, geri
    kalan 59 fıkra AYNI kalır.
    """
    document = docx_lib.Document()
    document.add_heading("BİRİNCİ BÖLÜM", level=1)
    document.add_heading("Uzun Madde", level=2)

    fikralar = []
    for i in range(1, 61):
        cumle = (
            f"Kurum, {i} numaralı süreçte ilgili birimlerle koordinasyonu "
            f"sürekli olarak sağlar ve yılda bir kez raporlar."
        )
        if degistir and i == 30:
            cumle = cumle.replace("sürekli olarak", "yılda en az bir kez")
        fikralar.append(f"({i}) {cumle}")

    document.add_paragraph(f"MADDE 1- {chr(10).join(fikralar)}")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestClassifyContentDiffTutarliligi:
    """
    Kalite kontrolünde bulunan YANLIŞ NEGATİF regresyonu: classify_content()
    TÜM madde metninin kaba benzerlik ORANINA bakar -- çok uzun bir maddede
    (örn. 60 fıkra) TEK bir cümlelik gerçek bir değişiklik, bu oranı İHMAL
    EDİLEBİLİR ölçüde etkiler ve eşiğin (IDENTICAL_CHAR_SIMILARITY_THRESHOLD)
    ÜSTÜNDE kalabilir -- madde GERÇEKTEN değişmiş olsa bile rozet "Değişmedi"
    derdi (section_diff aynı satırda bir "replace" cümlesi TAŞIDIĞI hâlde).
    compare_documents() artık bu iki sinyali BİRBİRİNE karşı doğruluyor
    (bkz. pipeline.py NEDEN notu).
    """

    def test_uzun_maddede_gomulu_tek_degisiklik_modified_olarak_isaretlenir(self):
        old_bytes = _build_docx_bytes(degistir=False)
        new_bytes = _build_docx_bytes(degistir=True)

        result = compare_documents(old_bytes, "uzun_eski.docx", new_bytes, "uzun_yeni.docx")

        assert len(result.rows) == 1
        row = result.rows[0]
        assert row.classified.change_type == ChangeType.MODIFIED
        # NEDEN section_diff'te TAM OLARAK bir "replace" bekleniyor: differ.py
        # ZATEN doğru çalışıyordu (bu regresyon SADECE classify_content'in
        # rozet kararındaydı) -- burada TEKRAR doğrulanıyor ki iki sinyal
        # ARTIK tutarlı: rozet MODIFIED derken diff de GERÇEKTEN tam olarak
        # bir fark buluyor (ne fazla ne eksik).
        replace_ops = [sd for sd in row.section_diff.sentence_diffs if sd.op == "replace"]
        assert len(replace_ops) == 1

    def test_hicbir_degisiklik_olmayan_uzun_madde_hala_identical(self):
        # NEDEN kritik: bu doğrulama YÖNÜ tek taraflı olmalı -- düzeltme
        # SADECE gerçekten farklı olan maddeleri MODIFIED'a çevirmeli,
        # gerçekten AYNI kalan uzun bir maddeyi YANLIŞLIKLA MODIFIED'a
        # ÇEVİRMEMELİ (bkz. pipeline.py NEDEN notu -- tek yönlü güvenlik ağı).
        same_bytes = _build_docx_bytes(degistir=False)

        result = compare_documents(same_bytes, "eski.docx", same_bytes, "yeni.docx")

        assert len(result.rows) == 1
        assert result.rows[0].classified.change_type == ChangeType.IDENTICAL


class TestMimariEk1MetaAktarimi:
    """
    Mimari Ek 1 (F1): compare_documents artık OPSİYONEL sürüm/tarih
    parametrelerini kabul edip old_meta/new_meta'ya (bkz. src/models.py::
    DocumentMeta) aktarıyor mu -- bu, F3/F4/F7'nin (belge_gorunen_adi,
    kaynak_atifi, render_comparison_direction) dayandığı verinin GERÇEKTEN
    boru hattından geçtiğini doğrular.
    """

    def test_parametre_verilmezse_slot_yine_de_atanir_digerleri_none(self):
        # NEDEN slot HER ZAMAN atanır: bkz. pipeline.py::compare_documents
        # NEDEN notu -- slot yükleme sırasına bağlıdır, tarih/etiket
        # girilip girilmediğinden BAĞIMSIZDIR.
        same_bytes = _build_docx_bytes(degistir=False)
        result = compare_documents(same_bytes, "eski.docx", same_bytes, "yeni.docx")

        assert result.old_meta.slot == "A"
        assert result.new_meta.slot == "B"
        assert result.old_meta.version_label is None
        assert result.old_meta.publication_date is None
        assert result.old_meta.effective_date is None

    def test_surum_ve_tarih_parametreleri_ilgili_meta_ya_aktarilir(self):
        same_bytes = _build_docx_bytes(degistir=False)
        result = compare_documents(
            same_bytes,
            "eski.docx",
            same_bytes,
            "yeni.docx",
            old_version_label="2019 sürümü",
            old_publication_date=date(2019, 11, 14),
            old_effective_date=date(2020, 1, 1),
            new_version_label="2023 sürümü",
            new_publication_date=date(2023, 6, 1),
            new_effective_date=date(2023, 9, 1),
        )

        assert result.old_meta.version_label == "2019 sürümü"
        assert result.old_meta.publication_date == date(2019, 11, 14)
        assert result.old_meta.effective_date == date(2020, 1, 1)
        assert result.new_meta.version_label == "2023 sürümü"
        assert result.new_meta.effective_date == date(2023, 9, 1)


class TestTemporalChangesEntegrasyonu:
    """
    ComparisonResult.temporal_changes'in GERÇEK bir DOCX boru hattından
    (ingestion -> structure_parser -> match_sections -> diff_section_match
    -> diff_temporal_expressions) uçtan uca DOĞRU DOLDUĞUNU doğrular --
    ayrıntılı eşleştirme/güven eşiği mantığı tests/test_temporal_diff.py'de
    zaten kapsanıyor, burada SADECE pipeline.py bağlantısı test edilir.
    """

    def test_sure_degisikligi_pipeline_uzerinden_tespit_edilir(self):
        old_bytes = _build_madde_docx_bytes("Başvurular 30 gün içinde değerlendirilir.")
        new_bytes = _build_madde_docx_bytes("Başvurular 45 gün içinde değerlendirilir.")

        result = compare_documents(old_bytes, "eski.docx", new_bytes, "yeni.docx")

        assert len(result.temporal_changes) == 1
        change = result.temporal_changes[0]
        assert change.paired is True
        assert change.old.text == "30 gün"
        assert change.new.text == "45 gün"

    def test_degisiklik_yoksa_bos_liste(self):
        same_bytes = _build_madde_docx_bytes("Bu maddenin amacı düzenleme yapmaktır.")

        result = compare_documents(same_bytes, "eski.docx", same_bytes, "yeni.docx")

        assert result.temporal_changes == []
