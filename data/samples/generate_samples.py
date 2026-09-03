"""
document_spec.py'deki TEK KAYNAK veriden dört test belgesi + ground_truth.json
üretir: yonetmelik_2019.pdf/.docx, yonetmelik_2023.pdf/.docx.

Çalıştırma: proje kökünden `python -m data.samples.generate_samples`.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from xml.sax.saxutils import escape

# NEDEN bu bootstrap: script doğrudan `python data/samples/generate_samples.py`
# olarak çalıştırılırsa Python sadece bu dosyanın klasörünü sys.path'e ekler;
# `src` ve `data` paketlerinin bulunabilmesi için proje kökü elle eklenir.
_PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

import docx  # noqa: E402
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.pdfbase import pdfmetrics  # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont  # noqa: E402
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer  # noqa: E402

from data.samples.document_spec import (  # noqa: E402
    KURUM_ADI,
    REGULATION_TITLE,
    get_article,
    get_bolumler,
    get_madde_no,
    resolved_text,
)
from src.config import SAMPLES_DIR  # noqa: E402

# --------------------------------------------------------------------------
# FONT ÇÖZÜMLEME
# --------------------------------------------------------------------------
# NEDEN: reportlab'in yerleşik (base14) Helvetica/Times fontları WinAnsi
# kodlamasıyla sınırlıdır ve ı/ğ/ş/İ gibi Türkçe'ye özgü karakterleri
# BASAMAZ. Gerçek bir TTF (tam Unicode kapsamı) kaydetmek şart.
_DEJAVU_CANDIDATES = [
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
    ),
    (
        "/usr/share/fonts/truetype/ttf-dejavu/DejaVuSerif.ttf",
        "/usr/share/fonts/truetype/ttf-dejavu/DejaVuSerif-Bold.ttf",
    ),
    (
        "/usr/local/share/fonts/DejaVuSerif.ttf",
        "/usr/local/share/fonts/DejaVuSerif-Bold.ttf",
    ),
]
# NEDEN ikinci sıra sistem Times TTF'i: DejaVuSerif kurulu değilse (örn. macOS),
# base14'e düşmeden ÖNCE gerçek bir Unicode TTF denenir -- bu, spesifikasyondaki
# "bulunamazsa Times'a düş" ifadesinin GERÇEK Türkçe karakter desteğiyle
# karşılanan hâlidir. Sadece bu da bulunamazsa base14 Times-Roman'a düşülür.
_SYSTEM_TTF_FALLBACK = (
    "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
    "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf",
)


def _register_body_fonts() -> tuple[str, str]:
    for regular_path, bold_path in _DEJAVU_CANDIDATES:
        if Path(regular_path).exists() and Path(bold_path).exists():
            pdfmetrics.registerFont(TTFont("BodyFont", regular_path))
            pdfmetrics.registerFont(TTFont("BodyFont-Bold", bold_path))
            return "BodyFont", "BodyFont-Bold"

    regular_path, bold_path = _SYSTEM_TTF_FALLBACK
    if Path(regular_path).exists() and Path(bold_path).exists():
        pdfmetrics.registerFont(TTFont("BodyFont", regular_path))
        pdfmetrics.registerFont(TTFont("BodyFont-Bold", bold_path))
        return "BodyFont", "BodyFont-Bold"

    print(
        "UYARI: DejaVuSerif veya sistem TTF bulunamadı; base14 Times-Roman "
        "kullanılacak. ı/ğ/ş/İ gibi Türkçe karakterler bozulabilir."
    )
    return "Times-Roman", "Times-Bold"


def _paragraph_markup(text: str) -> str:
    """reportlab Paragraph mini-XML'i için escape + \\n -> <br/> dönüşümü."""
    return escape(text).replace("\n", "<br/>\n")


def _turkish_upper(text: str) -> str:
    """
    Türkçe kurallarına uygun büyük harfe çevirme.

    NEDEN str.upper() YETERSİZ: Python'ın varsayılan Unicode büyütme kuralı
    küçük noktalı 'i' harfini noktasız 'I'ya çevirir (İngilizce kuralı).
    Türkçede noktalı 'i'nin büyüğü noktalı 'İ'dir -- bu satır olmadan
    üstbilgideki "Veri" kelimesi "VERI" olarak (yanlış) basılırdı.
    """
    return text.replace("i", "İ").upper()


# --------------------------------------------------------------------------
# PDF ÜRETİMİ
# --------------------------------------------------------------------------


def _make_page_decorator(font_name: str):
    """
    Her sayfaya tekrarlayan üstbilgi + 'Sayfa N' altbilgisi çizer.

    NEDEN sabit koordinatlar (üstten/alttan ~1cm): src/config.py'deki
    HEADER_BAND_RATIO/FOOTER_BAND_RATIO (%8) bandının rahatlıkla içinde
    kalır; SimpleDocTemplate margin'leri (3cm) ise gövde metninin bu banda
    hiç girmemesini garanti eder -- aksi halde pdf_loader'ın bant testi
    gövde metnini yanlışlıkla üstbilgi sanabilirdi.
    """

    def _decorate(canvas, doc) -> None:
        canvas.saveState()
        width, height = A4
        canvas.setFont(font_name, 8)
        canvas.drawCentredString(width / 2, height - 1.0 * cm, f"T.C. {_turkish_upper(KURUM_ADI)}")
        canvas.drawCentredString(width / 2, 1.0 * cm, f"Sayfa {canvas.getPageNumber()}")
        canvas.restoreState()

    return _decorate


def _build_pdf(year: int, output_path: Path, regular_font: str, bold_font: str) -> None:
    style_title = ParagraphStyle(
        "RegTitle", fontName=bold_font, fontSize=13, leading=17,
        alignment=TA_CENTER, spaceAfter=18,
    )
    style_bolum = ParagraphStyle(
        "Bolum", fontName=bold_font, fontSize=12, leading=15,
        alignment=TA_CENTER, spaceBefore=16, spaceAfter=2,
    )
    style_bolum_baslik = ParagraphStyle(
        "BolumBaslik", fontName=bold_font, fontSize=11, leading=14,
        alignment=TA_CENTER, spaceAfter=14,
    )
    style_madde_baslik = ParagraphStyle(
        "MaddeBaslik", fontName=bold_font, fontSize=10.5, leading=13,
        spaceBefore=8, spaceAfter=2,
    )
    style_madde_body = ParagraphStyle(
        "MaddeBody", fontName=regular_font, fontSize=10.5, leading=14,
        alignment=TA_JUSTIFY, spaceAfter=6,
    )

    flow: list = [Paragraph(escape(REGULATION_TITLE), style_title), Spacer(1, 6)]

    for bolum in get_bolumler(year):
        flow.append(Paragraph(f"{bolum.sira_sayisi} BÖLÜM", style_bolum))
        flow.append(Paragraph(escape(bolum.baslik), style_bolum_baslik))

        for article_id in bolum.article_ids:
            article = get_article(article_id)
            madde_no = get_madde_no(article, year)
            text = resolved_text(article, year)
            assert text is not None  # bu bölümdeki maddeler bu yılda var demektir

            flow.append(Paragraph(escape(article.baslik), style_madde_baslik))

            if article_id == "veri_kalitesi":
                # NEDEN kasıtlı PageBreak: "sayfa sınırında bölünen madde"
                # senaryosunu DETERMİNİSTİK üretmek için -- reportlab'in
                # otomatik akışına güvenmek yerine fıkra (1)'den sonra
                # zorla sayfa kesiyoruz; ingestion'ın bu maddeyi İKİ AYRI
                # bloğa (ve muhtemelen iki farklı sayfaya) böldüğünü,
                # Adım 2'nin bunları birleştirmesi gerektiğini test eder.
                fikralar = text.split("\n")
                ilk_fikra, kalan_fikralar = fikralar[0], "\n".join(fikralar[1:])
                flow.append(
                    Paragraph(f"MADDE {madde_no}- {_paragraph_markup(ilk_fikra)}", style_madde_body)
                )
                flow.append(PageBreak())
                flow.append(Paragraph(_paragraph_markup(kalan_fikralar), style_madde_body))
            else:
                flow.append(
                    Paragraph(f"MADDE {madde_no}- {_paragraph_markup(text)}", style_madde_body)
                )
            flow.append(Spacer(1, 6))

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        topMargin=3.0 * cm,
        bottomMargin=3.0 * cm,
        leftMargin=2.2 * cm,
        rightMargin=2.2 * cm,
        title=REGULATION_TITLE,
    )
    decorator = _make_page_decorator(regular_font)
    doc.build(flow, onFirstPage=decorator, onLaterPages=decorator)


# --------------------------------------------------------------------------
# DOCX ÜRETİMİ
# --------------------------------------------------------------------------


def _build_docx(year: int, output_path: Path) -> None:
    document = docx.Document()

    section = document.sections[0]
    # NEDEN is_linked_to_previous = False: aksi halde python-docx üstbilgi/
    # altbilgiyi "önceki bölümle aynı" kabul edip kendi içeriğimizi
    # kalıcı şekilde yazmayabilir.
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = f"T.C. {_turkish_upper(KURUM_ADI)}"
    section.footer.paragraphs[0].text = f"{KURUM_ADI} — Sayfa"

    # NEDEN level=0: python-docx'te level=0 "Title" stiline, level=1/2
    # "Heading 1"/"Heading 2" stillerine karşılık gelir -- bunlar GERÇEK
    # Word stilleridir (PDF'teki punto/kalınlık ipucunun aksine).
    document.add_heading(REGULATION_TITLE, level=0)

    for bolum in get_bolumler(year):
        document.add_heading(f"{bolum.sira_sayisi} BÖLÜM", level=1)
        document.add_heading(bolum.baslik, level=1)

        for article_id in bolum.article_ids:
            article = get_article(article_id)
            madde_no = get_madde_no(article, year)
            text = resolved_text(article, year)
            assert text is not None

            document.add_heading(article.baslik, level=2)
            # NEDEN add_paragraph(text) (add_run değil): python-docx, verilen
            # metindeki "\n" karakterlerini otomatik olarak gerçek Word satır
            # sonlarına (<w:br/>) çevirir; fıkra (1)/(2)/(3) ayrımı böylece
            # Word'de de görsel olarak korunur.
            document.add_paragraph(f"MADDE {madde_no}- {text}")

    document.save(str(output_path))


# --------------------------------------------------------------------------
# GROUND TRUTH ÜRETİMİ
# --------------------------------------------------------------------------


def _find_bolum_baslik(article_id: str, year: int) -> str:
    for bolum in get_bolumler(year):
        if article_id in bolum.article_ids:
            return f"{bolum.sira_sayisi} BÖLÜM"
    raise KeyError(f"'{article_id}' {year} yılı bölüm listelerinde bulunamadı")


def _build_ground_truth() -> dict:
    from data.samples.document_spec import ARTICLES

    maddeler = []
    for article in ARTICLES:
        bolum_2019 = (
            _find_bolum_baslik(article.id, 2019) if article.madde_no_2019 is not None else None
        )
        bolum_2023 = (
            _find_bolum_baslik(article.id, 2023) if article.madde_no_2023 is not None else None
        )
        maddeler.append(
            {
                "id": article.id,
                "change_type": article.change_type.value,
                "baslik": article.baslik,
                "madde_no_2019": article.madde_no_2019,
                "madde_no_2023": article.madde_no_2023,
                "bolum_2019": bolum_2019,
                "bolum_2023": bolum_2023,
                "aciklama": article.aciklama,
            }
        )

    return {
        "kaynak": "data/samples/document_spec.py",
        "yonetmelik_basligi": REGULATION_TITLE,
        "belgeler": {
            "2019": {"pdf": "yonetmelik_2019.pdf", "docx": "yonetmelik_2019.docx"},
            "2023": {"pdf": "yonetmelik_2023.pdf", "docx": "yonetmelik_2023.docx"},
        },
        "maddeler": maddeler,
    }


# --------------------------------------------------------------------------
# GİRİŞ NOKTASI
# --------------------------------------------------------------------------


def generate_all(output_dir: Path = SAMPLES_DIR) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    regular_font, bold_font = _register_body_fonts()

    _build_pdf(2019, output_dir / "yonetmelik_2019.pdf", regular_font, bold_font)
    _build_pdf(2023, output_dir / "yonetmelik_2023.pdf", regular_font, bold_font)
    _build_docx(2019, output_dir / "yonetmelik_2019.docx")
    _build_docx(2023, output_dir / "yonetmelik_2023.docx")

    ground_truth = _build_ground_truth()
    with open(output_dir / "ground_truth.json", "w", encoding="utf-8") as f:
        json.dump(ground_truth, f, ensure_ascii=False, indent=2)

    print(f"Üretildi: {output_dir}")


if __name__ == "__main__":
    generate_all()
