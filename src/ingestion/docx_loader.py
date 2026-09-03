"""python-docx tabanlı DOCX loader."""

from __future__ import annotations

from pathlib import Path

import docx

from src.config import MIN_TOTAL_CHARS_FOR_DOCX_TEXT_LAYER
from src.ingestion.base import DocumentLoader, NoTextLayerError
from src.models import Document, DocumentMeta, TextUnit


class DOCXLoader(DocumentLoader):
    """DOCX -> provenance'lı TextUnit listesi. Paragraf bazlı çıkarım kullanır."""

    def load(self, path: str | Path, doc_id: str) -> Document:
        path = Path(path)
        doc = docx.Document(str(path))

        # NEDEN `doc.paragraphs`: python-docx bu koleksiyonu SADECE gövde
        # (body) paragraflarından oluşturur; üstbilgi/altbilgi paragrafları
        # ayrı bir bölümde (section.header/section.footer) yaşar ve bu
        # koleksiyona hiç girmez. Yani "header/footer'ı gövdeye karıştırma"
        # kuralı, doğru koleksiyonu seçerek YAPISAL olarak sağlanır --
        # ayrıca bir filtreleme mantığına gerek yoktur.
        units: list[TextUnit] = []
        heading_style_hints: dict[int, str] = {}
        block_index = 0

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            # NEDEN "Heading"/"Title" ile başlayanları ipucu olarak saklıyoruz:
            # bunlar Word'ün YERLEŞİK başlık stilleridir ("Heading 1" .. "Heading 9",
            # "Title"); Adım 2'deki structure_parser bu ipucu, PDF'teki
            # punto/kalınlık ipucunun DOCX karşılığı olarak kullanacak.
            if style_name.startswith("Heading") or style_name == "Title":
                heading_style_hints[block_index] = style_name

            units.append(
                TextUnit(
                    doc_id=doc_id,
                    # NEDEN None: DOCX akış tabanlıdır, sayfa sınırları sadece
                    # render zamanında (yazı tipi/kağıt boyutuna göre) belirlenir;
                    # loader düzeyinde uydurma bir sayfa numarası GÜVENİLMEZ olur.
                    page_no=None,
                    block_index=block_index,
                    char_start=0,
                    char_end=len(text),
                    heading_path=(),
                    text=text,
                )
            )
            block_index += 1

        self._assert_text_layer_present(units, path)

        meta = DocumentMeta(
            doc_id=doc_id,
            source_path=str(path),
            file_type="docx",
            title=path.stem,
            page_count=None,
        )
        return Document(meta=meta, units=units, heading_style_hints=heading_style_hints)

    def _assert_text_layer_present(self, units: list[TextUnit], path: Path) -> None:
        total_chars = sum(len(u.text) for u in units)
        if total_chars < MIN_TOTAL_CHARS_FOR_DOCX_TEXT_LAYER:
            raise NoTextLayerError(
                f"'{path}' dosyasında yeterli metin bulunamadı "
                f"(toplam {total_chars} karakter). Belge boş veya sadece "
                "görüntü içeriyor olabilir."
            )
